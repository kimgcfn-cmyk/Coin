"""
최근 24시간 코인 투자 수익률 분석 엑셀 — 자동(봇)·수동 구분, 잔고·손익·현재 잔고, 종목별·봇별·합산 (2026-10-26)
================================================================================
사용:  python3 coin_pnl_report.py [--hours 24] [--out 폴더]     (맥에서 실행, ssh로 서버의 거래소 기록을 읽는다. 주문·변경 없음)
구분 규칙: 체결의 주문 ID가 봇 로그에 있으면 '자동', 없으면 '수동'(사용자 직접 거래 + 사용자 지시로 제가 낸 청산 주문).
잔고·손익: 계정 원장(financial records)의 잔고를 기준으로 24시간 전 지갑 잔고, 현재 지갑 잔고, 펀딩·수수료·실현손익을 맞춘다.
한계: 지갑은 하나라서 시작·현재 잔액은 합계로만 있다(자동/수동 잔액은 '시작 잔액 + 그 구분 손익' 가정치). 24시간 전 미실현 손익은 알 수 없어 0으로 둔다.
"""
import sys, os, json, subprocess, datetime as dt, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.utils import get_column_letter as L

KST = dt.timezone(dt.timedelta(hours=9))
OUT = "/Users/Muita/Desktop/Coin/코인봇수익률"
HERE = os.path.dirname(os.path.abspath(__file__))
# 사용자가 직접 연 물량을 봇 장부에 편입한 수량(봇이 관리하지만 개시는 수동): 미실현 배분에서 '수동 개시'로 본다
REGISTERED = {("ORDERUSDT", "short"): 2230.0, ("CAPUSDT", "long"): 4400.0}
BOTNAME = {("ORDERUSDT", "short"): "ORDER 숏 봇", ("TUSDT", "short"): "T 숏 봇", ("CAPUSDT", "long"): "CAP 롱 봇", ("CAPUSDT", "short"): "CAP 숏 봇", ("ARIAUSDT", "long"): "ARIA 롱 봇"}

def fetch(hours):
    src = open(os.path.join(HERE, "coin_pnl_fetch.py"), encoding="utf-8").read().replace("24*3600*1000", f"{hours}*3600*1000")
    r = subprocess.run(["ssh", "-o", "BatchMode=yes", "-o", "ConnectTimeout=15", "trading-ec2", "cd /home/ec2-user/trading-bot && venv/bin/python -"], input=src, capture_output=True, text=True, timeout=240)
    out = [l for l in r.stdout.splitlines() if l.startswith("{")]
    if not out: raise RuntimeError("서버 조회 실패: " + (r.stderr or r.stdout)[-300:])
    return json.loads(out[-1])

AR = "Arial"
f_norm = Font(name=AR, size=10); f_bold = Font(name=AR, size=10, bold=True); f_in = Font(name=AR, size=10, color="0000FF")
f_head = Font(name=AR, size=10, bold=True, color="FFFFFF"); f_title = Font(name=AR, size=14, bold=True); f_note = Font(name=AR, size=9, italic=True, color="595959")
f_green = Font(name=AR, size=10, bold=True, color="2E7D32")
fill_head = PatternFill("solid", fgColor="1F3864"); fill_sub = PatternFill("solid", fgColor="D9E2F3"); fill_key = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
NUM = "#,##0.00;[Red]-#,##0.00;-"; PCT = "0.00%;[Red]-0.00%;-"

def head(ws, r, labels, c0=1, h=32):
    for j, t in enumerate(labels, c0):
        c = ws.cell(r, j, t); c.font = f_head; c.fill = fill_head; c.alignment = ctr; c.border = box
    ws.row_dimensions[r].height = h

def build(d, path, hours, cum=None):
    ids = set(d["bot_ids"]); fills = sorted(d["fills"], key=lambda x: int(x["createdTime"])); recs = sorted(d["recs"], key=lambda x: int(x["ts"]))
    k = lambda ms: dt.datetime.fromtimestamp(int(ms) / 1000, KST).replace(tzinfo=None)
    t0 = k(d["start"]); t1 = k(d["now"])
    wb = Workbook()
    # ── 체결내역 ──
    tf = wb.active; tf.title = "체결내역"
    tf["A1"] = f"최근 {hours}시간 전체 체결 (거래소 기록)"; tf["A1"].font = f_title
    tf["A2"] = "구분: 주문 ID가 봇 로그에 있으면 자동, 없으면 수동. 거래소 실현손익은 수수료 전 값이고 순손익 = 실현손익 − 수수료."; tf["A2"].font = f_note
    H = ["No", "시각(KST)", "종목", "구분", "분류", "포지션", "거래", "체결가", "수량", "체결금액(USDT)", "수수료(USDT)", "실현손익(수수료 전)", "순손익(수수료 후)", "주문ID"]
    head(tf, 4, H); first = 5
    for i, x in enumerate(fills):
        r = first + i; sym = x["symbol"]; side = x["posSide"]; auto = x["orderId"] in ids
        fee = sum(float(y["fee"]) for y in x["feeDetail"])
        cls = BOTNAME.get((sym, side), "기타 자동") if auto else "수동 거래"
        trade = "진입" if str(x["tradeSide"]).startswith("open") else "청산"
        vals = [i + 1, k(x["createdTime"]), sym, "자동" if auto else "수동", cls, "롱" if side == "long" else "숏", trade, float(x["execPrice"]), float(x["execQty"]), float(x["execValue"]), fee, float(x["execPnl"]), None, x["orderId"]]
        for j, v in enumerate(vals, 1):
            c = tf.cell(r, j, v); c.font = f_in if j in (8, 9, 10, 11, 12) else f_norm; c.border = box
        tf.cell(r, 13, f"=L{r}-K{r}").font = f_norm; tf.cell(r, 13).border = box
        tf.cell(r, 2).number_format = "yyyy-mm-dd hh:mm:ss"; tf.cell(r, 8).number_format = "0.000000"; tf.cell(r, 9).number_format = "#,##0"
        for j in (10, 11, 12, 13): tf.cell(r, j).number_format = "#,##0.0000;[Red]-#,##0.0000;-"
        for j in (1, 4, 6, 7): tf.cell(r, j).alignment = Alignment(horizontal="center")
    last = first + len(fills) - 1
    for j, w in enumerate([6, 19, 14, 7, 14, 8, 7, 12, 11, 15, 13, 17, 17, 22], 1): tf.column_dimensions[L(j)].width = w
    tf.freeze_panes = tf.cell(first, 3)
    rg = lambda col: f"체결내역!${col}${first}:${col}${last}"
    # ── 잔고추이(원장) ──
    tb = wb.create_sheet("잔고추이")
    tb["A1"] = "계정 원장과 지갑 잔고 추이 (USDT, 실현 기준)"; tb["A1"].font = f_title
    tb["A2"] = "원장 1건 = 거래 수수료·실현손익·펀딩 정산 1건. 잔고(지갑)는 미실현 손익을 뺀 값입니다."; tb["A2"].font = f_note
    head(tb, 4, ["No", "시각(KST)", "유형", "종목", "금액(USDT)", "수수료", "잔고(지갑, USDT)"]); rf = 5
    for i, x in enumerate(recs):
        r = rf + i; vals = [i + 1, k(x["ts"]), x["type"], x.get("symbol") or "", float(x["amount"]), float(x.get("fee") or 0), float(x["balance"])]
        for j, v in enumerate(vals, 1):
            c = tb.cell(r, j, v); c.font = f_norm; c.border = box
        tb.cell(r, 2).number_format = "m/d hh:mm:ss"
        for j in (5, 6, 7): tb.cell(r, j).number_format = "#,##0.0000;[Red]-#,##0.0000;-"
    rl = rf + len(recs) - 1
    for j, w in enumerate([6, 17, 38, 14, 14, 11, 17], 1): tb.column_dimensions[L(j)].width = w
    tb.freeze_panes = "A5"
    rb = lambda col: f"잔고추이!${col}${rf}:${col}${rl}"
    # 시작·현재 잔고
    bal_start = (float(recs[0]["balance"]) - float(recs[0]["amount"])) if recs else 0.0
    bal_end = float(recs[-1]["balance"]) if recs else 0.0
    # ── 종합 요약 ──
    ws = wb.create_sheet("종합 요약", 0)
    ws["A1"] = f"코인 투자 {hours}시간 수익률 종합 — 자동(봇)·수동"; ws["A1"].font = f_title
    ws["A2"] = f"기간 {t0:%Y-%m-%d %H:%M} ~ {t1:%Y-%m-%d %H:%M} KST (거래소 기록 기준). 자동 = 봇이 낸 주문, 수동 = 사용자가 직접 낸 주문(및 사용자 지시로 제가 낸 청산)."; ws["A2"].font = f_note
    ws["A4"] = "① 계좌 잔고"; ws["A4"].font = f_bold
    rows = [("24시간 전 지갑 잔고 (원장 기준)", bal_start, "시작 시점 직후 첫 원장 기록의 잔고에서 그 기록 금액을 뺀 값"),
            ("기간 내 입출금·이체 순액", 0.0, "원장에 입출금·이체 유형이 없어 0 (유형: 거래·수수료·펀딩만)"),
            ("현재 지갑 잔고 (원장 마지막 기록)", bal_end, "미실현 제외"),
            ("현재 미실현 손익", float(d["assets"]["unrealisedPnl"]), "열려 있는 포지션 평가"),
            ("현재 총자산 (지갑 + 미실현)", float(d["assets"]["accountEquity"]), "거래소 계정 총자산"),
            ("현재 사용 증거금", float(d["assets"]["imr"]), "")]
    for i, (lab, v, note) in enumerate(rows):
        r = 5 + i; ws.cell(r, 1, lab).font = f_norm; c = ws.cell(r, 2, v); c.font = f_in; c.number_format = NUM; c.border = box; ws.cell(r, 3, note).font = f_note
    ws["B5"].fill = fill_key; ws["B7"].fill = fill_key
    ws["A12"] = "② 자동·수동 손익 (24시간)"; ws["A12"].font = f_bold
    head(ws, 13, ["구분", "체결 수", "거래대금(USDT)", "실현손익(수수료 전)", "거래 수수료", "실현 순손익(수수료 후)", "현재 미실현(추정 배분)", "총 손익(실현+미실현)", "시작 잔고 대비 수익률", "시작 잔액 + 이 구분 손익(가정)"], h=42)
    # 미실현 배분(수량 비례)
    upl = {"자동": 0.0, "수동": 0.0}; pos_rows = []
    for p in d["positions"]:
        sym, side = p["symbol"], p["side"]; reg = REGISTERED.get((sym, side), 0.0)
        # 봇 장부 수량(상태 파일)로 봇이 관리하는 수량 파악: ORDER/T/CAP
        nm = {("ORDERUSDT", "short"): "ORDER", ("TUSDT", "short"): "T", ("CAPUSDT", "long"): "CAP", ("CAPUSDT", "short"): "CAP_S", ("ARIAUSDT", "long"): "ARIA"}.get((sym, side)); st = d["states"].get(nm, {}) if nm else {}
        bot_qty = float(st.get("qty") or 0.0) if (BOTNAME.get((sym, side)) and nm) else 0.0
        bot_opened = max(0.0, bot_qty - reg); manual_opened = max(0.0, p["qty"] - bot_opened)
        share_auto = (bot_opened / p["qty"]) if p["qty"] else 0.0
        upl["자동"] += p["upl"] * share_auto; upl["수동"] += p["upl"] * (1 - share_auto)
        pos_rows.append((sym, "롱" if side == "long" else "숏", p["qty"], p["avg"], p["mark"], p["notional"], p["upl"], bot_opened, manual_opened, p["upl"] * share_auto, p["upl"] * (1 - share_auto)))
    for i, g_ in enumerate(("자동", "수동")):
        r = 14 + i
        ws.cell(r, 1, g_).font = f_bold
        ws.cell(r, 2, f'=COUNTIFS({rg("D")},"{g_}")'); ws.cell(r, 3, f'=SUMIFS({rg("J")},{rg("D")},"{g_}")')
        ws.cell(r, 4, f'=SUMIFS({rg("L")},{rg("D")},"{g_}")'); ws.cell(r, 5, f'=SUMIFS({rg("K")},{rg("D")},"{g_}")')
        ws.cell(r, 6, f"=D{r}-E{r}"); ws.cell(r, 7, upl[g_]); ws.cell(r, 8, f"=F{r}+G{r}"); ws.cell(r, 9, f"=H{r}/$B$5"); ws.cell(r, 10, f"=$B$5+H{r}")
    r = 16; ws.cell(r, 1, "펀딩비 (공통, 구분 배분 불가)").font = f_norm
    ws.cell(r, 6, f'=SUMIFS({rb("E")},{rb("C")},"*SETTLE_FEE*")'); ws.cell(r, 8, f"=F{r}"); ws.cell(r, 9, f"=H{r}/$B$5")
    r = 17; ws.cell(r, 1, "합계").font = f_bold
    for col in "BCDEFGH": ws[f"{col}{r}"] = f"=SUM({col}14:{col}16)"
    ws[f"I{r}"] = f"=H{r}/$B$5"; ws[f"J{r}"] = f"=$B$5+H{r}"
    for rr in range(14, 18):
        for cc in range(1, 11):
            c = ws.cell(rr, cc); c.border = box; c.font = f_bold if (rr == 17 or cc == 1) else f_norm
            if cc in (3, 4, 5, 6, 7, 8, 10): c.number_format = NUM
            if cc == 9: c.number_format = PCT
            if cc == 2: c.number_format = "#,##0"
        if rr == 17:
            for cc in range(1, 11): ws.cell(rr, cc).fill = fill_sub
    ws["A19"] = "③ 검증 (거래소 원장과 맞는지)"; ws["A19"].font = f_bold
    chk = [("시작 잔고 + 실현 순손익 + 펀딩 (체결 기준)", "=B5+B6+F17", "체결의 실현손익·수수료와 펀딩을 더한 값"),
           ("현재 지갑 잔고 (원장)", "=B7", ""),
           ("차이", "=B20-B21", "0에 가까워야 함 (미세 차이는 반올림·원장 시각 경계)"),
           ("현재 총자산 − 시작 잔고 (입출금 제외)", "=B9-B5-B6", "실현 + 현재 미실현 − (시작 미실현은 알 수 없어 0 가정)"),
           ("총 손익 합계 (위 표)", "=H17", ""),
           ("계정 총자산 변화율 (시작 잔고 대비)", "=B23/B5", "")]
    for i, (lab, f, note) in enumerate(chk):
        rr = 20 + i; ws.cell(rr, 1, lab).font = f_norm; c = ws.cell(rr, 2, f); c.font = f_bold if i in (2, 5) else f_norm; c.number_format = PCT if i == 5 else NUM; c.border = box; ws.cell(rr, 3, note).font = f_note
    ws["A27"] = "※ 시작·현재 잔액은 지갑이 하나라 합계로만 존재합니다. 표의 '시작 잔액 + 이 구분 손익'은 각 구분의 손익만 반영했을 때의 가상 잔액입니다(합쳐도 실제 잔액이 되지 않음). 미실현 손익은 포지션 수량 비례로 배분한 추정치입니다."; ws["A27"].font = f_note
    ws["A28"] = "※ 과거 결과일 뿐 투자 조언이 아닙니다."; ws["A28"].font = f_note
    for j, w in enumerate([46, 15, 17, 20, 15, 22, 21, 20, 18, 24], 1): ws.column_dimensions[L(j)].width = w
    # ── 종목별 ──
    wy = wb.create_sheet("종목별", 1)
    wy["A1"] = "종목별·구분별 손익 (24시간)"; wy["A1"].font = f_title
    head(wy, 3, ["종목", "구분", "체결 수", "거래대금(USDT)", "실현손익(수수료 전)", "수수료", "실현 순손익", "펀딩"], h=34)
    syms = sorted({x["symbol"] for x in fills}); r = 4
    for sym in syms:
        for g_ in ("자동", "수동"):
            if not any(x["symbol"] == sym and (x["orderId"] in ids) == (g_ == "자동") for x in fills): continue
            wy.cell(r, 1, sym); wy.cell(r, 2, g_)
            crit = f'{rg("C")},A{r},{rg("D")},B{r}'
            wy.cell(r, 3, f"=COUNTIFS({crit})"); wy.cell(r, 4, f"=SUMIFS({rg('J')},{crit})"); wy.cell(r, 5, f"=SUMIFS({rg('L')},{crit})"); wy.cell(r, 6, f"=SUMIFS({rg('K')},{crit})"); wy.cell(r, 7, f"=E{r}-F{r}")
            for cc in range(1, 9): wy.cell(r, cc).border = box; wy.cell(r, cc).font = f_norm
            for cc in (4, 5, 6, 7): wy.cell(r, cc).number_format = NUM
            r += 1
        fr = r - 1
    # 종목별 펀딩 (원장)
    wy.cell(r + 1, 1, "종목별 펀딩 (원장)").font = f_bold; head(wy, r + 2, ["종목", "펀딩(USDT)"]); rr = r + 3
    for sym in syms:
        wy.cell(rr, 1, sym).font = f_norm; c = wy.cell(rr, 2, f'=SUMIFS({rb("E")},{rb("C")},"*SETTLE_FEE*",{rb("D")},A{rr})'); c.number_format = NUM; c.font = f_norm
        wy.cell(rr, 1).border = box; c.border = box; rr += 1
    rr += 1; wy.cell(rr, 1, "현재 열린 포지션과 미실현 배분").font = f_bold; head(wy, rr + 1, ["종목", "방향", "수량", "평균가", "현재가", "명목(USDT)", "미실현(USDT)", "봇 개시 수량", "수동 개시 수량", "자동 배분", "수동 배분"], h=34); q = rr + 2
    for row in pos_rows:
        for j, v in enumerate(row, 1):
            c = wy.cell(q, j, v); c.font = f_norm; c.border = box
            if j in (3, 8, 9): c.number_format = "#,##0"
            if j in (4, 5): c.number_format = "0.000000"
            if j in (6, 7, 10, 11): c.number_format = NUM
        q += 1
    wy.cell(q + 1, 1, "※ 봇 개시 수량 = 봇 장부 수량 − 사용자가 직접 연 뒤 봇에 편입한 수량(ORDER 숏 2,230개, CAP 롱 4,400개). 미실현 배분은 수량 비례 추정입니다.").font = f_note
    for j, w in enumerate([16, 8, 12, 17, 19, 12, 16, 14, 14, 13, 13], 1): wy.column_dimensions[L(j)].width = w
    # ── 봇별 ──
    wbk = wb.create_sheet("봇별(자동)", 2)
    wbk["A1"] = "자동(봇)별 손익 (24시간)"; wbk["A1"].font = f_title
    head(wbk, 3, ["봇", "체결 수", "진입 체결", "청산 체결", "거래대금(USDT)", "실현손익(수수료 전)", "수수료", "실현 순손익", "청산 1건당 평균 순손익"], h=34)
    bots = [n for n in ("ORDER 숏 봇", "T 숏 봇", "CAP 롱 봇", "CAP 숏 봇", "ARIA 롱 봇", "기타 자동") if any(((BOTNAME.get((x["symbol"], x["posSide"]), "기타 자동") if x["orderId"] in ids else None) == n) for x in fills)]
    for i, n in enumerate(bots):
        r = 4 + i; wbk.cell(r, 1, n); crit = f'{rg("E")},A{r}'
        wbk.cell(r, 2, f"=COUNTIFS({crit})"); wbk.cell(r, 3, f'=COUNTIFS({crit},{rg("G")},"진입")'); wbk.cell(r, 4, f'=COUNTIFS({crit},{rg("G")},"청산")')
        wbk.cell(r, 5, f"=SUMIFS({rg('J')},{crit})"); wbk.cell(r, 6, f"=SUMIFS({rg('L')},{crit})"); wbk.cell(r, 7, f"=SUMIFS({rg('K')},{crit})"); wbk.cell(r, 8, f"=F{r}-G{r}"); wbk.cell(r, 9, f"=IF(D{r}=0,0,H{r}/D{r})")
        for cc in range(1, 10):
            wbk.cell(r, cc).border = box; wbk.cell(r, cc).font = f_norm
            if cc in (5, 6, 7, 8, 9): wbk.cell(r, cc).number_format = NUM
    r = 4 + len(bots); wbk.cell(r, 1, "합계").font = f_bold
    for cc in range(2, 9):
        c = wbk.cell(r, cc, f"=SUM({L(cc)}4:{L(cc)}{r-1})"); c.font = f_bold; c.border = box; c.number_format = NUM if cc >= 5 else "#,##0"
    wbk.cell(r, 1).border = box
    wbk.cell(r + 2, 1, "※ 봇별 손익은 거래소 체결(평균가 처리는 거래소 기준) 합계라서 봇 장부의 평균단가법 손익과 소폭 다를 수 있습니다. 진입 체결의 실현손익은 0입니다.").font = f_note
    for j, w in enumerate([16, 10, 11, 11, 17, 19, 12, 15, 21], 1): wbk.column_dimensions[L(j)].width = w
    # ── 잔고 차트 (잔고추이 시트) ──
    ch = ScatterChart(); ch.title = "지갑 잔고 추이 (USDT, 실현 기준)"; ch.style = 2
    ch.x_axis.title = "시각 (KST)"; ch.y_axis.title = "지갑 잔고 (USDT)"
    se = Series(Reference(tb, min_col=7, min_row=rf, max_row=rl), Reference(tb, min_col=2, min_row=rf, max_row=rl), title="지갑 잔고")
    se.marker.symbol = "none"; se.smooth = False; se.graphicalProperties.line.width = 22000; se.graphicalProperties.line.solidFill = "2E7D32"
    ch.series.append(se); ch.legend = None; ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.x_axis.number_format = "m/d hh:mm"; ch.y_axis.number_format = "#,##0"; ch.x_axis.tickLblPos = "low"; ch.y_axis.tickLblPos = "low"
    for ttl in (ch.title, ch.x_axis.title, ch.y_axis.title): ttl.overlay = False
    ch.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge", x=0.11, y=0.14, w=0.85, h=0.62))
    for ax in (ch.x_axis, ch.y_axis): ax.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9", w=6350))
    ch.height = 10; ch.width = 24; tb.add_chart(ch, "I4")
    # ── 설명·가정 ──
    wd = wb.create_sheet("설명·가정")
    wd["A1"] = "설명·가정"; wd["A1"].font = f_title
    notes = ["· 데이터: Bitget 선물 계정의 체결 기록(실현손익 포함)과 계정 원장(잔고·수수료·펀딩). 기간은 종합 요약 시트 위쪽에 표시된 24시간입니다.",
             "· 자동/수동: 체결의 주문 ID가 봇 실행 로그에 있으면 '자동', 없으면 '수동'입니다. 사용자가 직접 낸 거래와, 사용자 지시로 제가 낸 청산 주문(예: T 숏 전량 청산)은 '수동'에 들어갑니다.",
             "· 편입 물량: 사용자가 직접 연 뒤 봇 장부에 편입한 물량(ORDER 숏 2,230개, CAP 롱 4,400개)은 체결 기준으로는 수동, 편입 이후 봇이 관리합니다. 편입 이후 봇이 낸 매도·환매는 '자동'으로 분류됩니다.",
             "· 실현손익(수수료 전): 거래소가 체결마다 계산한 값입니다. 순손익 = 실현손익 − 수수료. 펀딩비는 원장의 정산 기록이며 종목 단위로만 알 수 있어 자동/수동에는 배분하지 않았습니다.",
             "· 시작 잔고: 24시간 시점 직후 첫 원장 기록의 잔고에서 그 금액을 뺀 값입니다. 지갑 잔고는 미실현 손익을 포함하지 않으므로, 24시간 전 미실현 손익은 0으로 가정했습니다.",
             "· 미실현 배분: 열려 있는 포지션의 미실현을 봇이 개시한 수량과 수동 개시 수량 비율로 나눈 추정치입니다.",
             "· 수익률: 손익 ÷ 24시간 전 지갑 잔고(입출금 없음). 자동·수동 수익률은 그 구분이 계좌 전체에 기여한 비율입니다.",
             "· 값 칸(파란 글자)은 거래소에서 가져온 원자료이고, 검은 글자는 수식입니다. 과거 결과일 뿐 투자 조언이 아닙니다."]
    for i, t in enumerate(notes): wd.cell(3 + i, 1, t).font = f_norm
    wd.column_dimensions["A"].width = 160
    if cum: add_cumulative(wb, cum)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return dict(fills=len(fills), recs=len(recs), bal_start=bal_start, bal_end=bal_end, upl=upl)

def add_cumulative(wb, D):
    """자동매매 누적 현황(코인별)·종합분석 시트 (2026-10-26). D = 긴 구간(96시간) 조회 결과. 자동 체결만 누적 성과에 쓰고, 수동은 비교용."""
    ids = set(D["bot_ids"]); fills = sorted(D["fills"], key=lambda x: int(x["createdTime"]))
    k = lambda ms: dt.datetime.fromtimestamp(int(ms) / 1000, KST).replace(tzinfo=None)
    _auto_ts = [int(x["createdTime"]) for x in fills if x["orderId"] in ids]
    if _auto_ts:                                                   # 같은 기간 비교를 위해 첫 자동 체결 이후만 쓴다(그 전 수동 거래는 제외)
        _t0 = min(_auto_ts); fills = [x for x in fills if int(x["createdTime"]) >= _t0]; D = dict(D, recs=[x for x in D["recs"] if int(x["ts"]) >= _t0])
    # ── 누적 체결내역 ──
    tf = wb.create_sheet("누적체결내역")
    tf["A1"] = "누적 체결 (첫 자동 체결 이후, 자동·수동 전체)"; tf["A1"].font = f_title
    tf["A2"] = "자동 = 주문 ID가 봇 로그에 있는 체결. 누적 자동 순손익·낙폭은 자동 체결만 더한 값입니다(진입 체결은 수수료만 반영)."; tf["A2"].font = f_note
    head(tf, 4, ["No", "시각(KST)", "종목", "구분", "분류", "포지션", "거래", "체결가", "수량", "체결금액(USDT)", "수수료(USDT)", "실현손익(수수료 전)", "순손익(수수료 후)", "누적 자동 순손익", "누적 최고", "낙폭"])
    first = 5; autos = []
    for i, x in enumerate(fills):
        r = first + i; sym = x["symbol"]; side = x["posSide"]; auto = x["orderId"] in ids; fee = sum(float(y["fee"]) for y in x["feeDetail"])
        cls = BOTNAME.get((sym, side), "기타 자동") if auto else "수동 거래"; trade = "진입" if str(x["tradeSide"]).startswith("open") else "청산"
        pnl = float(x["execPnl"])
        vals = [i + 1, k(x["createdTime"]), sym, "자동" if auto else "수동", cls, "롱" if side == "long" else "숏", trade, float(x["execPrice"]), float(x["execQty"]), float(x["execValue"]), fee, pnl]
        for j, v in enumerate(vals, 1):
            c = tf.cell(r, j, v); c.font = f_in if j in (8, 9, 10, 11, 12) else f_norm; c.border = box
        tf.cell(r, 13, f"=L{r}-K{r}"); prev_n = f"N{r-1}" if i else "0"; prev_o = f"O{r-1}" if i else "0"
        tf.cell(r, 14, f'={prev_n}+IF(D{r}="자동",M{r},0)'); tf.cell(r, 15, f"=MAX({prev_o},N{r})"); tf.cell(r, 16, f"=N{r}-O{r}")
        for j in range(13, 17): tf.cell(r, j).font = f_norm; tf.cell(r, j).border = box
        tf.cell(r, 2).number_format = "yyyy-mm-dd hh:mm:ss"; tf.cell(r, 8).number_format = "0.000000"; tf.cell(r, 9).number_format = "#,##0"
        for j in (10, 11, 12, 13, 14, 15, 16): tf.cell(r, j).number_format = "#,##0.0000;[Red]-#,##0.0000;-"
        for j in (1, 4, 6, 7): tf.cell(r, j).alignment = Alignment(horizontal="center")
        if auto: autos.append(dict(t=k(x["createdTime"]), sym=sym, side=side, trade=trade, net=pnl - fee, pnl=pnl, fee=fee, amt=float(x["execValue"])))
    last = first + len(fills) - 1
    for j, w in enumerate([6, 19, 14, 7, 14, 8, 7, 12, 11, 15, 13, 17, 17, 17, 14, 12], 1): tf.column_dimensions[L(j)].width = w
    tf.freeze_panes = tf.cell(first, 3)
    rg = lambda col: f"누적체결내역!${col}${first}:${col}${last}"
    # ── 파이썬 쪽 요약(코멘트용) ──
    syms = sorted({a["sym"] for a in autos}, key=lambda s: -sum(a["net"] for a in autos if a["sym"] == s))
    cum = peak = dd = 0.0
    for a in autos:
        cum += a["net"]; peak = max(peak, cum); dd = min(dd, cum - peak)
    closes = [a for a in autos if a["trade"] == "청산"]; wins = [a["net"] for a in closes if a["net"] > 0]; losses = [a["net"] for a in closes if a["net"] < 0]
    net_all = sum(a["net"] for a in autos); gross = sum(a["pnl"] for a in autos); fee_all = sum(a["fee"] for a in autos); vol_all = sum(a["amt"] for a in autos)
    man = [x for x in fills if x["orderId"] not in ids]
    man_net = sum(float(x["execPnl"]) - sum(float(y["fee"]) for y in x["feeDetail"]) for x in man)
    t_first = min(a["t"] for a in autos) if autos else None; t_last = max(a["t"] for a in autos) if autos else None
    span_h = ((t_last - t_first).total_seconds() / 3600) if autos else 0.0
    fund = sum(float(x["amount"]) for x in D["recs"] if "SETTLE_FEE" in str(x.get("type")))
    # ── 코인별 시트 ──
    wc = wb.create_sheet("자동매매누적", 3)
    wc["A1"] = "자동매매 누적 현황 — 코인별"; wc["A1"].font = f_title
    wc["A2"] = f"기간 {t_first:%m/%d %H:%M} ~ {t_last:%m/%d %H:%M} KST, 자동(봇) 체결만 집계 (거래소 체결 기준)." if autos else "자동 체결이 없습니다."; wc["A2"].font = f_note
    head(wc, 4, ["종목", "봇", "첫 자동 체결", "마지막 자동 체결", "가동 시간(h)", "자동 체결 수", "청산 체결 수", "거래대금(USDT)", "실현손익(수수료 전)", "수수료", "실현 순손익", "청산 1건당 평균", "청산 승률", "최대 단건 이익", "최대 단건 손실", "현재 보유(거래소)", "현재 미실현", "총 손익(실현+미실현)", "시간당 순손익", "봇 장부 누적(상태 파일)"], h=36)
    stmap = {("ORDERUSDT", "short"): "ORDER", ("TUSDT", "short"): "T", ("CAPUSDT", "long"): "CAP", ("CAPUSDT", "short"): "CAP_S", ("ARIAUSDT", "long"): "ARIA"}
    r0 = 5; n = len(syms)
    for i, s in enumerate(syms):
        r = r0 + i; mine = [a for a in autos if a["sym"] == s]; cl = [a["net"] for a in mine if a["trade"] == "청산"]
        bots = "·".join(sorted({BOTNAME.get((s, a["side"]), "기타") for a in mine}))
        pos = [p for p in D["positions"] if p["symbol"] == s]; pos_txt = " / ".join(f"{'롱' if p['side']=='long' else '숏'} {p['qty']:,.0f}개({p['notional']:,.0f}USDT)" for p in pos) or "없음"
        upl = sum(p["upl"] for p in pos)
        stv = sum(float((D["states"].get(nm) or {}).get("cum_pnl") or 0.0) for (sy, sd), nm in stmap.items() if sy == s)
        vals = [s, bots, min(a["t"] for a in mine), max(a["t"] for a in mine)]
        for j, v in enumerate(vals, 1): c = wc.cell(r, j, v); c.font = f_bold if j == 1 else f_norm; c.border = box
        crit = f'{rg("C")},$A{r},{rg("D")},"자동"'
        wc.cell(r, 5, f"=(D{r}-C{r})*24"); wc.cell(r, 6, f"=COUNTIFS({crit})"); wc.cell(r, 7, f'=COUNTIFS({crit},{rg("G")},"청산")')
        wc.cell(r, 8, f"=SUMIFS({rg('J')},{crit})"); wc.cell(r, 9, f"=SUMIFS({rg('L')},{crit})"); wc.cell(r, 10, f"=SUMIFS({rg('K')},{crit})"); wc.cell(r, 11, f"=I{r}-J{r}")
        wc.cell(r, 12, f"=IF(G{r}=0,0,K{r}/G{r})"); wc.cell(r, 13, f'=IF(G{r}=0,0,COUNTIFS({crit},{rg("G")},"청산",{rg("M")},">0")/G{r})')
        wc.cell(r, 14, max([x for x in cl if x > 0] or [0.0])); wc.cell(r, 15, min([x for x in cl if x < 0] or [0.0])); wc.cell(r, 16, pos_txt); wc.cell(r, 17, upl)
        wc.cell(r, 18, f"=K{r}+Q{r}"); wc.cell(r, 19, f"=IF(E{r}<=0,0,K{r}/E{r})"); wc.cell(r, 20, stv)
        for j in range(5, 21):
            c = wc.cell(r, j); c.border = box; c.font = f_in if j in (14, 15, 17, 20) else f_norm
            if j in (8, 9, 10, 11, 12, 14, 15, 17, 18, 19, 20): c.number_format = NUM
            if j == 5: c.number_format = "0.0"
            if j == 13: c.number_format = "0%"
            if j in (6, 7): c.number_format = "#,##0"
        wc.cell(r, 3).number_format = "m/d hh:mm"; wc.cell(r, 4).number_format = "m/d hh:mm"
    rt = r0 + n; wc.cell(rt, 1, "합계").font = f_bold
    for j in (6, 7, 8, 9, 10, 11, 17, 18, 20):
        c = wc.cell(rt, j, f"=SUM({L(j)}{r0}:{L(j)}{rt-1})"); c.font = f_bold; c.number_format = NUM if j >= 8 else "#,##0"
    c = wc.cell(rt, 12, f"=IF(G{rt}=0,0,K{rt}/G{rt})"); c.number_format = NUM; c.font = f_bold
    c = wc.cell(rt, 13, f'=IF(G{rt}=0,0,COUNTIFS({rg("D")},"자동",{rg("G")},"청산",{rg("M")},">0")/G{rt})'); c.number_format = "0%"; c.font = f_bold
    for j in range(1, 21): wc.cell(rt, j).border = box; wc.cell(rt, j).fill = fill_sub
    wc.cell(rt + 1, 1, "※ 순손익 = 거래소 실현손익 − 수수료(진입 체결은 수수료만 반영). 최대 단건 이익·손실은 청산 체결 1건 기준. 자동으로 분류된 체결에는 봇 장부에 편입한 수동 물량의 환매·청산이 포함됩니다.").font = f_note
    wc.cell(rt + 2, 1, "※ 봇 장부 누적(상태 파일)은 봇이 평균단가법으로 센 값이라 거래소 체결 합계와 다를 수 있습니다(편입 시 평균가 재설정·수동 거래 반영 차이).").font = f_note
    # 봇별 표
    b0 = rt + 4; wc.cell(b0 - 1, 1, "봇별 누적").font = f_bold
    head(wc, b0, ["봇", "상태", "첫 자동 체결", "마지막 자동 체결", "가동 시간(h)", "체결 수", "청산 체결 수", "거래대금(USDT)", "실현손익(수수료 전)", "수수료", "실현 순손익", "청산 1건당 평균", "청산 승률"], h=34)
    botnames = [nm for nm in ("ORDER 숏 봇", "T 숏 봇", "CAP 롱 봇", "CAP 숏 봇", "ARIA 롱 봇") if any(BOTNAME.get((a["sym"], a["side"])) == nm for a in autos)]
    for i, nm in enumerate(botnames):
        r = b0 + 1 + i; mine = [a for a in autos if BOTNAME.get((a["sym"], a["side"])) == nm]
        for j, v in enumerate([nm, "가동 중" if (D.get("active") or {}).get(nm) else "중지", min(a["t"] for a in mine), max(a["t"] for a in mine)], 1):
            c = wc.cell(r, j, v); c.font = f_bold if j == 1 else f_norm; c.border = box
        crit = f'{rg("E")},$A{r},{rg("D")},"자동"'
        wc.cell(r, 5, f"=(D{r}-C{r})*24"); wc.cell(r, 6, f"=COUNTIFS({crit})"); wc.cell(r, 7, f'=COUNTIFS({crit},{rg("G")},"청산")')
        wc.cell(r, 8, f"=SUMIFS({rg('J')},{crit})"); wc.cell(r, 9, f"=SUMIFS({rg('L')},{crit})"); wc.cell(r, 10, f"=SUMIFS({rg('K')},{crit})"); wc.cell(r, 11, f"=I{r}-J{r}")
        wc.cell(r, 12, f"=IF(G{r}=0,0,K{r}/G{r})"); wc.cell(r, 13, f'=IF(G{r}=0,0,COUNTIFS({crit},{rg("G")},"청산",{rg("M")},">0")/G{r})')
        for j in range(5, 14):
            c = wc.cell(r, j); c.border = box; c.font = f_norm
            if j in (8, 9, 10, 11, 12): c.number_format = NUM
            if j == 5: c.number_format = "0.0"
            if j == 13: c.number_format = "0%"
            if j in (6, 7): c.number_format = "#,##0"
        wc.cell(r, 3).number_format = "m/d hh:mm"; wc.cell(r, 4).number_format = "m/d hh:mm"
    for j, w in enumerate([14, 20, 13, 13, 11, 11, 11, 15, 17, 11, 14, 15, 10, 14, 14, 30, 13, 17, 14, 20], 1): wc.column_dimensions[L(j)].width = w
    wc.freeze_panes = "B5"
    # ── 종합분석 시트 ──
    wz = wb.create_sheet("자동매매종합", 1)
    wz["A1"] = "자동매매 종합 분석 (누적)"; wz["A1"].font = f_title
    wz["A2"] = f"자동(봇) 체결 누적 기준 — 기간 {t_first:%m/%d %H:%M} ~ {t_last:%m/%d %H:%M} KST. 숫자는 수식, 아래 평가 문장은 생성 시점 값으로 만든 요약입니다." if autos else "자동 체결이 없습니다."; wz["A2"].font = f_note
    A = rg("D"); G = rg("G"); M = rg("M")
    items = [("① 기간·규모", None, None),
             ("첫 자동 체결", t_first, "yyyy-mm-dd hh:mm"), ("마지막 자동 체결", t_last, "yyyy-mm-dd hh:mm"), ("경과 시간(h)", "=(B7-B6)*24", "0.0"),
             ("자동 체결 수", f'=COUNTIFS({A},"자동")', "#,##0"), ("거래대금(USDT)", f'=SUMIFS({rg("J")},{A},"자동")', NUM), ("거래 수수료(USDT)", f'=SUMIFS({rg("K")},{A},"자동")', NUM), ("수수료 ÷ 거래대금", "=B11/B10", "0.000%"),
             ("② 성과", None, None),
             ("실현손익(수수료 전)", f'=SUMIFS({rg("L")},{A},"자동")', NUM), ("실현 순손익(수수료 후)", "=B14-B11", NUM), ("수수료가 수수료 전 이익에서 차지한 비중", "=IF(B14=0,0,B11/B14)", "0.0%"),
             ("청산 체결 수", f'=COUNTIFS({A},"자동",{G},"청산")', "#,##0"), ("청산 승률(순손익>0)", f'=IF(B17=0,0,COUNTIFS({A},"자동",{G},"청산",{M},">0")/B17)', "0.0%"),
             ("평균 이익 (승리 청산)", f'=IF(COUNTIFS({A},"자동",{G},"청산",{M},">0")=0,0,SUMIFS({M},{A},"자동",{G},"청산",{M},">0")/COUNTIFS({A},"자동",{G},"청산",{M},">0"))', NUM),
             ("평균 손실 (패배 청산)", f'=IF(COUNTIFS({A},"자동",{G},"청산",{M},"<0")=0,0,SUMIFS({M},{A},"자동",{G},"청산",{M},"<0")/COUNTIFS({A},"자동",{G},"청산",{M},"<0"))', NUM),
             ("손익비 (평균 이익 ÷ 평균 손실 절대값)", "=IF(B20=0,0,B19/-B20)", "0.00"), ("청산 1건당 평균 순손익", "=IF(B17=0,0,B15/B17)", NUM), ("시간당 순손익", "=IF(B8=0,0,B15/B8)", NUM),
             ("최대 낙폭 (누적 자동 순손익 기준)", f'=MIN({rg("P")})', NUM), ("현재 미실현(자동분 추정)", None, NUM), ("총 손익(실현 + 미실현)", "=B15+B25", NUM)]
    auto_upl = 0.0
    for p in D["positions"]:
        nm = stmap.get((p["symbol"], p["side"])); st = D["states"].get(nm, {}) if nm else {}
        if nm and p["qty"]: auto_upl += p["upl"] * min(1.0, float(st.get("qty") or 0.0) / p["qty"])
    rr = 4
    for lab, v, fmt in items:
        rr += 1
        if v is None and fmt is None and lab.startswith(("①", "②")): wz.cell(rr, 1, lab).font = f_bold; continue
        wz.cell(rr, 1, lab).font = f_norm
        c = wz.cell(rr, 2, auto_upl if lab.startswith("현재 미실현") else v); c.number_format = fmt; c.border = box; c.font = f_in if (lab.startswith("현재 미실현") or isinstance(v, dt.datetime)) else f_norm
    # 셀 위치: ① B5..B11 → rows 5..; 위 수식 참조와 맞는지 확인용으로 행 번호를 고정한다
    wz["A28"] = "③ 자동 vs 수동 (같은 기간)"; wz["A28"].font = f_bold
    head(wz, 29, ["구분", "체결 수", "거래대금(USDT)", "실현손익(수수료 전)", "수수료", "실현 순손익"], h=30)
    for i, g_ in enumerate(("자동", "수동")):
        r = 30 + i; wz.cell(r, 1, g_).font = f_bold
        wz.cell(r, 2, f'=COUNTIFS({A},"{g_}")'); wz.cell(r, 3, f'=SUMIFS({rg("J")},{A},"{g_}")'); wz.cell(r, 4, f'=SUMIFS({rg("L")},{A},"{g_}")'); wz.cell(r, 5, f'=SUMIFS({rg("K")},{A},"{g_}")'); wz.cell(r, 6, f"=D{r}-E{r}")
    wz.cell(32, 1, "펀딩비 (공통)").font = f_norm; c = wz.cell(32, 6, fund); c.font = f_in
    wz.cell(33, 1, "합계").font = f_bold
    for j in range(2, 7): wz.cell(33, j, f"=SUM({L(j)}30:{L(j)}32)")
    for rr_ in range(30, 34):
        for j in range(1, 7):
            c = wz.cell(rr_, j); c.border = box
            if j > 1 and c.font != f_in: c.font = f_bold if rr_ == 33 else f_norm
            if j == 2: c.number_format = "#,##0"
            elif j > 2: c.number_format = NUM
    # ④ 평가 문장
    best = syms[0] if syms else "-"; worst = syms[-1] if syms else "-"; nets = {s: sum(a["net"] for a in autos if a["sym"] == s) for s in syms}
    wr = (len(wins) / len(closes)) if closes else 0.0; aw = (sum(wins) / len(wins)) if wins else 0.0; al = (sum(losses) / len(losses)) if losses else 0.0
    pf = (aw / -al) if al else 0.0
    notes = [f"1) 규모·수익: 자동매매는 {span_h:.0f}시간 동안 {len(autos)}건 체결(거래대금 {vol_all:,.0f} USDT)했고 실현 순손익은 {net_all:+.1f} USDT입니다. 수수료 {fee_all:.1f} USDT가 수수료 전 이익({gross:+.1f})의 {(fee_all / gross if gross else 0):.0%}를 차지했습니다.",
             f"2) 코인별: " + " / ".join(f"{s.replace('USDT', '')} {nets[s]:+.1f}" for s in syms) + f" — {best.replace('USDT', '')}가 가장 좋았고 {worst.replace('USDT', '')}가 가장 낮았습니다.",
             f"3) 승률·손익비: 청산 {len(closes)}건 중 승률 {wr:.0%}, 평균 이익 {aw:+.2f}·평균 손실 {al:+.2f}, 손익비 {pf:.2f}. " + ("작은 이익을 자주 거두는 구조(손익비 1 미만)라 드물게 큰 손실이 나면 이익이 한꺼번에 줄어듭니다." if (pf and pf < 1 and wr > 0.6) else ("이익이 손실보다 큰 구조입니다." if pf >= 1 else "승률·손익비를 더 지켜볼 필요가 있습니다.")),
             f"4) 위험: 누적 자동 순손익 기준 최대 낙폭은 {dd:.1f} USDT, 시간당 순손익은 {(net_all / span_h if span_h else 0):+.2f} USDT입니다. 그리드는 한 방향으로 쏠리는 장(상승 중 숏, 하락 중 롱)에서 손실이 쌓이는 구조라 한도·손절선이 결과를 좌우합니다.",
             f"5) 현재 가동: " + (", ".join(n for n, a in (D.get("active") or {}).items() if a) or "확인 불가") + ". 중지: " + (", ".join(n for n, a in (D.get("active") or {}).items() if not a) or "없음") + ".",
             f"6) 수동 비교: 같은 기간 수동 거래의 순손익은 {man_net:+.1f} USDT(체결 {len(man)}건)로, 자동 {net_all:+.1f}와 합치면 {net_all + man_net:+.1f} USDT(펀딩 {fund:+.1f} 별도)입니다. 방향을 직접 바꾼 구간의 손익이 결과를 크게 좌우했습니다." if man else "6) 수동 체결이 없습니다.",
             "7) 해석 시 유의: 자동 체결에는 봇 장부에 편입한 수동 물량의 환매·청산이 포함되고, 편입 전 수동 개시 손익은 수동으로 분류됩니다. 과거 결과일 뿐 투자 조언이 아닙니다."]
    wz["A35"] = "④ 종합 평가"; wz["A35"].font = f_bold
    for i, t in enumerate(notes): c = wz.cell(36 + i, 1, t); c.font = f_norm; c.alignment = Alignment(wrap_text=True, vertical="top"); wz.merge_cells(start_row=36 + i, start_column=1, end_row=36 + i, end_column=6); wz.row_dimensions[36 + i].height = 48
    for j, w in enumerate([44, 20, 18, 20, 14, 18], 1): wz.column_dimensions[L(j)].width = w
    # 누적 자동 순손익 차트
    if autos:
        ch = ScatterChart(); ch.title = "누적 자동 순손익 (USDT)"; ch.style = 2; ch.x_axis.title = "시각 (KST)"; ch.y_axis.title = "누적 순손익 (USDT)"
        se = Series(Reference(tf, min_col=14, min_row=first, max_row=last), Reference(tf, min_col=2, min_row=first, max_row=last), title="누적 자동 순손익")
        se.marker.symbol = "none"; se.smooth = False; se.graphicalProperties.line.width = 22000; se.graphicalProperties.line.solidFill = "2E7D32"
        ch.series.append(se); ch.legend = None; ch.x_axis.delete = False; ch.y_axis.delete = False; ch.x_axis.number_format = "m/d hh:mm"; ch.y_axis.number_format = "#,##0"
        ch.x_axis.tickLblPos = "low"; ch.y_axis.tickLblPos = "low"
        for ttl in (ch.title, ch.x_axis.title, ch.y_axis.title): ttl.overlay = False
        ch.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge", x=0.11, y=0.14, w=0.85, h=0.62))
        for ax in (ch.x_axis, ch.y_axis): ax.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9", w=6350))
        ch.height = 9; ch.width = 22; wz.add_chart(ch, "H4")


def main():
    hours = 24; out = OUT; cum_hours = 96
    if "--hours" in sys.argv: hours = int(sys.argv[sys.argv.index("--hours") + 1])
    if "--out" in sys.argv: out = sys.argv[sys.argv.index("--out") + 1]
    os.makedirs(out, exist_ok=True)
    full = fetch(max(cum_hours, hours))                      # 누적 분석용으로 길게 한 번 조회하고
    s24 = full["now"] - hours * 3600 * 1000                   # 24시간 표는 그 구간만 잘라 쓴다
    d = dict(full, start=s24, fills=[x for x in full["fills"] if int(x["createdTime"]) >= s24], recs=[x for x in full["recs"] if int(x["ts"]) >= s24 - 1])
    path = os.path.join(out, f"코인투자_{hours}시간_수익률분석.xlsx")
    tmp = path + ".tmp.xlsx"; info = build(d, tmp, hours, cum=full); os.replace(tmp, path); print("저장:", path, info)   # 임시 파일에 만든 뒤 교체(자동 갱신 중 엑셀이 반쯤 쓴 파일을 읽지 않게)

if __name__ == "__main__":
    main()
