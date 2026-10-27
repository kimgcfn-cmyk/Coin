import re, sys
from openpyxl import load_workbook
def split_args(s):
    out=[]; depth=0; q=False; cur=""
    for ch in s:
        if ch=='"': q=not q
        if not q:
            if ch=="(": depth+=1
            elif ch==")": depth-=1
            elif ch=="," and depth==0: out.append(cur); cur=""; continue
        cur+=ch
    out.append(cur); return out
def find_close(s,i):  # s[i]=='('
    d=0; q=False
    for j in range(i,len(s)):
        ch=s[j]
        if ch=='"': q=not q
        if q: continue
        if ch=="(": d+=1
        elif ch==")":
            d-=1
            if d==0: return j
def tr(e):
    # IF 재귀 변환
    while True:
        m=re.search(r"\bIF\(",e)
        if not m: break
        i=m.end()-1; j=find_close(e,i); args=split_args(e[i+1:j])
        rep="((%s) if (%s) else (%s))"%(tr(args[1]),tr(args[0]),tr(args[2]))
        e=e[:m.start()]+"\0"+rep.replace("(","\1").replace(")","\2")+"\0"+e[j+1:]
    e=e.replace("\1","(").replace("\2",")").replace("\0","")
    return e
def to_py(f):
    e=f[1:]
    import re as _r2
    e=_r2.sub(r'"([<>=]*)"&\(([^()]*)\)', r'"\1"+str(\2)', e)
    e=tr(e)
    # 범위 먼저: 시트!$A$1:$A$9
    e=re.sub(r"(?:'([^']+)'|([^\s,()+\-*/=<>!']+))!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)",lambda m:'RNG("%s","%s",%s,%s)'%(m.group(1) or m.group(2),m.group(3),m.group(4),m.group(6)),e)
    e=re.sub(r"(?:'([^']+)'|([^\s,()+\-*/=<>!']+))!\$?([A-Z]+)\$?(\d+)",lambda m:'C("%s","%s%s")'%(m.group(1) or m.group(2),m.group(3),m.group(4)),e)
    # 같은 시트 참조 (따옴표 문자열 밖)
    parts=re.split(r'("[^"]*")',e); out=[]
    for p in parts:
        if p.startswith('"'): out.append(p); continue
        p=re.sub(r"(?<![A-Za-z\"_])\$?([A-Z]{1,2})\$?(\d+)(?![\d(])",lambda m:'C(SH,"%s%s")'%(m.group(1),m.group(2)),p)
        p=re.sub(r"(?<![<>=!])=(?!=)","==",p); p=p.replace("<>","!=")
        out.append(p)
    res="".join(out)
    res=_r2.sub(r'C\(SH,"([A-Z]+)(\d+)"\):C\(SH,"([A-Z]+)(\d+)"\)', r'RNG(SH,"\1",\2,\4)', res)
    return res
import fnmatch as _fn, re as _re3
def _match(x, c):
    if isinstance(c,(int,float)): return isinstance(x,(int,float)) and x==c
    c=str(c); mm=_re3.match(r"^(>=|<=|<>|>|<|=)?(.*)$",c); op=mm.group(1) or "="; v=mm.group(2)
    try: fv=float(v); isnum=True
    except ValueError: isnum=False
    if isnum and op in (">=","<=",">","<"): return isinstance(x,(int,float)) and {">":x>fv,">=":x>=fv,"<":x<fv,"<=":x<=fv}[op]
    if isnum and op=="=" and isinstance(x,(int,float)): return x==fv
    return _fn.fnmatchcase(str(x),v) if op in ("=",) else (not _fn.fnmatchcase(str(x),v))
def SUMIFSX(sr,*a):
    pairs=list(zip(a[0::2],a[1::2])); tot=0.0
    for i,v in enumerate(sr):
        if isinstance(v,(int,float)) and all(_match(rg[i],c) for rg,c in pairs): tot+=v
    return tot
def COUNTIFSX(*a):
    pairs=list(zip(a[0::2],a[1::2])); n=len(pairs[0][0]); return sum(1 for i in range(n) if all(_match(rg[i],c) for rg,c in pairs))
def run(path):
    wb=load_workbook(path); cache={}
    def C(sh,ref):
        k=(sh,ref)
        if k in cache: return cache[k]
        v=wb[sh][ref].value
        if v is None: v=""
        if isinstance(v,str) and v.startswith("="):
            SH=sh
            import datetime as _d
            def DATEVALUE(x): return (_d.date.fromisoformat(x)-_d.date(1899,12,30)).days
            def TIMEVALUE(x): hh,mm=x.split(':'); return (int(hh)*60+int(mm))/1440
            env=dict(DATEVALUE=DATEVALUE,TIMEVALUE=TIMEVALUE,LEFT=lambda x,n:str(x)[:n],MID=lambda x,a,n:str(x)[a-1:a-1+n],C=C,SH=sh,RNG=RNG,N=N,ROUND=lambda x,n:round(x,n),SUM=SUM,COUNTIF=COUNTIF,COUNTIFS=COUNTIFSX,SUMIFS=SUMIFSX,LOOKUP=LOOKUP,IFERROR=lambda f,d:d)
            expr=to_py(v)
            try:
                if "IFERROR" in expr: raise RuntimeError
                v=eval(expr,env)
            except RuntimeError:
                m=re.match(r"IFERROR\((.*),0\)$",v[1:]); inner="="+m.group(1)
                try: v=eval(to_py(inner),env)
                except Exception: v=0
        cache[k]=v; return v
    def RNG(sh,col,a,b): return [C(sh,f"{col}{r}") for r in range(a,b+1)]
    def N(x): return x if isinstance(x,(int,float)) and not isinstance(x,bool) else 0
    def SUM(l): return sum(x for x in l if isinstance(x,(int,float)))
    def COUNTIF(l,v): return sum(1 for x in l if x==v)
    def COUNTIFS(l1,v1,l2,v2):
        op=v2[0]; th=float(v2[1:]); return sum(1 for a_,b_ in zip(l1,l2) if a_==v1 and isinstance(b_,(int,float)) and b_>th)
    def _crit(c):
        import re as _r
        m=_r.match(r"(>=|<=|>|<|=)?(.*)",c); op=m.group(1) or "="; v=float(m.group(2)); return op,v
    def _ok(x,op,v):
        if not isinstance(x,(int,float)): return False
        return {">":x>v,">=":x>=v,"<":x<v,"<=":x<=v,"=":x==v}[op]
    def SUMIFS(s,cr,c):
        op,v=_crit(c); return sum(a for a,b_ in zip(s,cr) if isinstance(a,(int,float)) and _ok(b_,op,v))
    def COUNTIFS_1(cr,c):
        op,v=_crit(c); return sum(1 for b_ in cr if _ok(b_,op,v))
    def LOOKUP(big,l):
        nums=[x for x in l if isinstance(x,(int,float))]; return nums[-1]
    return wb,C
if __name__=="__main__":
    path=sys.argv[1]; wb,C=run(path)
    t=wb["거래내역"]; last=t.max_row
    n=sum(1 for r in range(5,last+1) if t.cell(r,5).value not in (None,""))
    print(path.split("/")[-1],"체결",n)
    s=wb["요약"]
    for r in range(4,26): print(f"  {s.cell(r,1).value:<34} {C('요약',f'B{r}')}")
