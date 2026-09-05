"""
거래 이력 자동 기록 — 매수/매도 체결마다 CSV 원장에 기록하고,
곧바로 코인매매_수익률.xlsx를 재생성한다.

coin_range_monitor.py(BTC), coin_grid_monitor.py(TQQQ)에서
체결 성공(if ok:) 시점에 append_trade()를 호출해서 쓴다.

CSV를 원장(source of truth)으로 두고 엑셀은 매번 통째로 재생성하는 이유:
거래가 하루 몇 건 수준이라 성능 문제가 없고, 재생성 방식이 부분수정보다
훨씬 덜 깨진다(엑셀 파일 중간에 행 삽입/누적식 갱신하다 꼬이는 사고 방지).
"""
import csv
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

LEDGER_DIR = os.path.dirname(os.path.abspath(__file__))
XLSX_PATH = os.path.join(LEDGER_DIR, "코인매매_수익률.xlsx")
SEED_TOTAL = 900.0   # 시드머니 총액 (2026-09-05 사용자 확인)

BOTS = {
    "BTC":  {"ledger": "btc_ledger.csv",  "leverage": 3, "margin": 10.0},
    "TQQQ": {"ledger": "tqqq_ledger.csv", "leverage": 1, "margin": 10.0},
}

FIELDS = ["buy_time", "buy_price", "sell_time", "sell_price", "note"]

def _ledger_path(bot: str) -> str:
    return os.path.join(LEDGER_DIR, BOTS[bot]["ledger"])

def _load_rows(path: str) -> list:
    if not os.path.exists(path):
        return []
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))

def _save_rows(path: str, rows: list):
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in FIELDS})

def append_trade(bot: str, side: str, time_str: str, price: float, note: str = ""):
    """
    side: "buy" | "sell"
    매수: 새 미청산(open) 행 추가.
    매도: 가장 최근 미청산 행을 찾아 매도가/매도시각을 채워 닫는다.
    실패해도(예: 엑셀 잠김) 원장 CSV 자체는 반드시 남긴다 — 엑셀 재생성은
    별개로 try/except 처리.
    """
    path = _ledger_path(bot)
    rows = _load_rows(path)

    if side == "buy":
        rows.append({"buy_time": time_str, "buy_price": price,
                      "sell_time": "", "sell_price": "", "note": note})
    elif side == "sell":
        closed = False
        for r in reversed(rows):
            if r.get("sell_time", "") == "":
                r["sell_time"] = time_str
                r["sell_price"] = price
                if note:
                    r["note"] = (r.get("note", "") + " " + note).strip()
                closed = True
                break
        if not closed:
            # 매수 기록 없이 매도만 들어온 이례적 케이스 — 그래도 남긴다
            rows.append({"buy_time": "", "buy_price": "",
                          "sell_time": time_str, "sell_price": price,
                          "note": ("⚠️매수기록없음 " + note).strip()})
    else:
        raise ValueError(f"알 수 없는 side: {side}")

    _save_rows(path, rows)

    try:
        rebuild_excel()
    except Exception as e:
        print(f"  ⚠️ 엑셀 리포트 갱신 실패(CSV 원장 저장은 정상 완료): {e}")

# ══════════════════════════════════════════════════════
#  엑셀 재생성
# ══════════════════════════════════════════════════════
HEADER = ["번호", "매수일시", "매수가", "매도일시", "매도가", "상태",
          "가격수익률(%)", "레버리지반영수익률(%)", "투입증거금(USDT)",
          "실현손익(USDT)", "누적손익(USDT)", "누적수익률(%)", "비고"]

HDR_FILL = PatternFill("solid", fgColor="305496")
HDR_FONT = Font(color="FFFFFF", bold=True)
OPEN_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")

def _write_sheet(ws, bot: str):
    cfg = BOTS[bot]
    rows = _load_rows(_ledger_path(bot))

    for c, h in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font, cell.alignment, cell.border = HDR_FILL, HDR_FONT, CENTER, BORDER

    r_out = 2
    cum_pnl = 0.0
    realized_pnl = 0.0
    open_pnl = 0.0
    for i, r in enumerate(rows, 1):
        buy_price = float(r["buy_price"]) if r["buy_price"] else None
        is_open = (r.get("sell_time", "") == "")

        if is_open:
            vals = [i, r["buy_time"], buy_price, "(보유중)", "", "보유중(미실현)",
                    "-", "-", cfg["margin"], "-", "-", "-", r.get("note", "")]
            fill = OPEN_FILL
        else:
            sell_price = float(r["sell_price"])
            price_ret = (sell_price - buy_price) / buy_price * 100 if buy_price else 0
            lev_ret = price_ret * cfg["leverage"]
            pnl = cfg["margin"] * lev_ret / 100
            cum_pnl += pnl
            realized_pnl = cum_pnl
            cum_ret = cum_pnl / SEED_TOTAL * 100
            vals = [i, r["buy_time"], buy_price, r["sell_time"], sell_price, "완료",
                    round(price_ret, 2), round(lev_ret, 2), cfg["margin"],
                    round(pnl, 3), round(cum_pnl, 3), round(cum_ret, 4), r.get("note", "")]
            fill = None

        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r_out, column=c, value=v)
            cell.border = BORDER
            cell.alignment = CENTER
            if fill:
                cell.fill = fill
        r_out += 1

    widths = [6, 17, 12, 17, 12, 16, 13, 17, 15, 13, 13, 13, 20]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    return realized_pnl

def rebuild_excel():
    wb = Workbook()
    ws_btc = wb.active
    ws_btc.title = "BTC 거래이력"
    ws_tqqq = wb.create_sheet("TQQQ 거래이력")
    ws_sum = wb.create_sheet("요약")

    btc_realized = _write_sheet(ws_btc, "BTC")
    tqqq_realized = _write_sheet(ws_tqqq, "TQQQ")
    total_realized = btc_realized + tqqq_realized

    def _closed_count(bot):
        return sum(1 for r in _load_rows(_ledger_path(bot)) if r.get("sell_time"))

    summary = [
        ("시드머니(USDT)", SEED_TOTAL),
        ("", ""),
        ("BTC 완료거래 수", _closed_count("BTC")),
        ("BTC 실현손익(USDT)", round(btc_realized, 3)),
        ("", ""),
        ("TQQQ 완료거래 수", _closed_count("TQQQ")),
        ("TQQQ 실현손익(USDT)", round(tqqq_realized, 3)),
        ("", ""),
        ("전체 실현손익 합계(USDT)", round(total_realized, 3)),
        ("전체 실현 누적수익률(%)", round(total_realized / SEED_TOTAL * 100, 4)),
        ("", ""),
        ("마지막 갱신", __import__("datetime").datetime.now(
            __import__("datetime").timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
    ]
    for r, (label, val) in enumerate(summary, 1):
        lc = ws_sum.cell(row=r, column=1, value=label)
        vc = ws_sum.cell(row=r, column=2, value=val)
        if label:
            lc.font = Font(bold=True)
        lc.border = BORDER
        vc.border = BORDER
    ws_sum.column_dimensions["A"].width = 30
    ws_sum.column_dimensions["B"].width = 20

    wb.save(XLSX_PATH)
